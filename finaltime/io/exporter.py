"""Экспорт отчётов в Excel и CSV."""

import csv
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from finaltime import get_version
from finaltime.core import AgeCategory, FinishRecord, format_time_short


class Exporter:
    """Экспорт результатов в Excel и CSV."""

    def __init__(self, logger):
        self.logger = logger

    def export_excel(
        self,
        finishes: list[FinishRecord],
        output_dir: str,
        registration_file: str = "",
        competition_date: date | None = None,
        categories: list[AgeCategory] | None = None,
        competition_name: str = "",
    ) -> str:
        """Экспорт в красивый Excel-файл.

        Структура отчёта:
        1. Заголовок (название соревнований, дата, файл регистрации)
        2. ОБЩИЙ ЗАЧЁТ — все участники, отсортированные по итоговому времени
        3. Таблицы по категориям — по одной на каждую категорию из файла
           категорий (в том же порядке), участники отфильтрованы и
           отсортированы, места внутри категории.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Результаты"

        styles = {
            'header_font': Font(name='Calibri', bold=True, size=12, color='FFFFFF'),
            'header_fill': PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid'),
            'header_alignment': Alignment(horizontal='center', vertical='center', wrap_text=True),
            'title_font': Font(name='Calibri', bold=True, size=16, color='2C3E50'),
            'subtitle_font': Font(name='Calibri', size=11, color='555555'),
            'section_font': Font(name='Calibri', bold=True, size=14, color='FFFFFF'),
            'section_fill': PatternFill(start_color='34495E', end_color='34495E', fill_type='solid'),
            'data_font': Font(name='Calibri', size=11),
            'data_alignment': Alignment(horizontal='center', vertical='center'),
            'name_alignment': Alignment(horizontal='left', vertical='center'),
            'thin_border': Border(
                left=Side(style='thin', color='CCCCCC'),
                right=Side(style='thin', color='CCCCCC'),
                top=Side(style='thin', color='CCCCCC'),
                bottom=Side(style='thin', color='CCCCCC'),
            ),
            'gold_fill': PatternFill(start_color='FFD700', end_color='FFD700', fill_type='solid'),
            'silver_fill': PatternFill(start_color='C0C0C0', end_color='C0C0C0', fill_type='solid'),
            'bronze_fill': PatternFill(start_color='CD7F32', end_color='CD7F32', fill_type='solid'),
        }

        headers = [
            "Место", "Номер", "ФИО участника", "Разряд", "Категория",
            "Время таймера", "Ручное время",
            "Штраф (сек)", "Штраф (баллы)", "Итоговое время"
        ]
        ncols = len(headers)

        # === ЗАГОЛОВОК ОТЧЁТА ===
        last_col = get_column_letter(ncols)
        ws.merge_cells(f'A1:{last_col}1')
        if competition_name.strip():
            ws['A1'] = competition_name.strip().upper()
        else:
            ws['A1'] = "ПРОТОКОЛ РЕЗУЛЬТАТОВ СОРЕВНОВАНИЙ"
        ws['A1'].font = styles['title_font']
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 35

        ws.merge_cells(f'A2:{last_col}2')
        ws['A2'] = f"Сгенерировано FinalTime v{get_version()}"
        ws['A2'].font = styles['subtitle_font']
        ws['A2'].alignment = Alignment(horizontal='center')

        if competition_date:
            ws.merge_cells(f'A3:{last_col}3')
            ws['A3'] = f"Дата соревнований: {competition_date.strftime('%d.%m.%Y')}"
            ws['A3'].font = styles['subtitle_font']
            ws['A3'].alignment = Alignment(horizontal='center')
            row = 5
        else:
            row = 4

        if registration_file:
            ws.merge_cells(f'A{row}:{last_col}{row}')
            ws[f'A{row}'] = f"Файл регистрации: {Path(registration_file).name}"
            ws[f'A{row}'].font = styles['subtitle_font']
            ws[f'A{row}'].alignment = Alignment(horizontal='center')
            row += 1

        # === ОБЩИЙ ЗАЧЁТ ===
        overall = [(r, r.place) for r in finishes]
        overall_header_row = row + 1  # строка заголовков колонок общего зачёта
        row = self._write_result_table(
            ws, row, overall, headers, competition_date, categories, styles,
            section_title="ОБЩИЙ ЗАЧЁТ",
        )

        # === ТАБЛИЦЫ ПО КАТЕГОРИЯМ (в порядке файла категорий) ===
        if competition_date and categories:
            for cat in categories:
                filtered = [
                    r for r in finishes
                    if r.participant.date_of_birth
                    and cat.matches(r.participant.gender, r.participant.age_on_date(competition_date))
                ]
                if not filtered:
                    continue
                cat_sorted = sorted(
                    filtered,
                    key=lambda r: r.final_time if r.final_time is not None else timedelta.max,
                )
                records_with_places = [(r, i) for i, r in enumerate(cat_sorted, 1)]
                row = self._write_result_table(
                    ws, row + 1, records_with_places, headers,
                    competition_date, categories, styles,
                    section_title=cat.name,
                )

        # === НАСТРОЙКА ШИРИНЫ КОЛОНОК ===
        column_widths = [8, 10, 35, 10, 12, 15, 15, 12, 14, 16]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # Заморозка верхних строк (заголовок + общий зачёт)
        ws.freeze_panes = f'A{overall_header_row + 1}'

        # Автофильтр по общему зачёту
        ws.auto_filter.ref = f'A{overall_header_row}:{last_col}{overall_header_row + len(finishes)}'

        # Сохранение
        output_path = Path(output_dir) / self._build_report_filename(competition_name, "xlsx")
        wb.save(output_path)

        self.logger.log("INFO", f"Excel report saved: {output_path}")
        return str(output_path)

    def _write_result_table(
        self,
        ws,
        start_row: int,
        records_with_places: list,
        headers: list,
        competition_date: date | None,
        categories: list[AgeCategory] | None,
        styles: dict,
        section_title: str,
    ) -> int:
        """Пишет секцию отчёта: заголовок секции + таблицу с данными.

        records_with_places: список кортежей (record, place).
        Возвращает номер следующей свободной строки.
        """
        ncols = len(headers)
        last_col = get_column_letter(ncols)

        # Заголовок секции
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=ncols)
        cell = ws.cell(row=start_row, column=1, value=section_title)
        cell.font = styles['section_font']
        cell.fill = styles['section_fill']
        cell.alignment = Alignment(horizontal='center', vertical='center')
        for col_idx in range(1, ncols + 1):
            ws.cell(row=start_row, column=col_idx).border = styles['thin_border']
        ws.row_dimensions[start_row].height = 24
        start_row += 1

        # Заголовки колонок
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = styles['header_font']
            cell.fill = styles['header_fill']
            cell.alignment = styles['header_alignment']
            cell.border = styles['thin_border']
        ws.row_dimensions[start_row].height = 30

        # Данные
        data_start = start_row + 1
        for offset, (record, place) in enumerate(records_with_places):
            row_idx = data_start + offset

            # Места 1-3 — подсветка
            row_fill = None
            if place == 1:
                row_fill = styles['gold_fill']
            elif place == 2:
                row_fill = styles['silver_fill']
            elif place == 3:
                row_fill = styles['bronze_fill']

            timer_str = format_time_short(record.timer_time)
            manual_str = format_time_short(record.manual_time) if record.manual_time else ""
            final_str = format_time_short(record.final_time) if record.final_time else ""

            # Категория
            category_str = "-"
            if competition_date and categories and record.participant.date_of_birth:
                category_str = record.participant.category_label(competition_date, categories)

            row_data = [
                place,
                record.participant.number,
                record.participant.full_name,
                record.participant.rank,
                category_str,
                timer_str,
                manual_str,
                record.penalty_seconds,
                record.penalty_points,
                final_str,
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = styles['data_font']
                cell.border = styles['thin_border']
                if col_idx == 3:  # ФИО
                    cell.alignment = styles['name_alignment']
                else:
                    cell.alignment = styles['data_alignment']
                if row_fill:
                    cell.fill = row_fill

        return data_start + len(records_with_places)

    def _build_report_filename(self, competition_name: str, ext: str) -> str:
        """Формирует имя файла отчёта: <название>_<дата>_<время>.<ext>."""
        name = competition_name.strip()
        if not name:
            name = "FinalTime"
        # Заменяем недопустимые символы и пробелы
        import re
        name = re.sub(r'[\\/:*?"<>|]', '_', name)
        name = re.sub(r'\s+', '_', name).strip('_')
        # Дата и время создания файла
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        return f"{name}_{timestamp}.{ext}"

    def export_csv(
        self,
        finishes: list[FinishRecord],
        output_dir: str,
        competition_date: date | None = None,
        categories: list[AgeCategory] | None = None,
        competition_name: str = "",
    ) -> str:
        """Export to CSV."""
        output_path = Path(output_dir) / self._build_report_filename(competition_name, "csv")

        with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f, delimiter=';')

            # Header
            writer.writerow([
                "Place", "Number", "Name", "Rank", "Category",
                "Timer Time", "Manual Time",
                "Penalty (sec)", "Penalty (pts)", "Final Time"
            ])

            # Data
            for record in finishes:
                timer_str = format_time_short(record.timer_time)
                manual_str = format_time_short(record.manual_time) if record.manual_time else ""
                final_str = format_time_short(record.final_time) if record.final_time else ""

                category_str = "-"
                if competition_date and categories and record.participant.date_of_birth:
                    category_str = record.participant.category_label(competition_date, categories)

                writer.writerow([
                    record.place,
                    record.participant.number,
                    record.participant.full_name,
                    record.participant.rank,
                    category_str,
                    timer_str,
                    manual_str,
                    record.penalty_seconds,
                    record.penalty_points,
                    final_str,
                ])

        self.logger.log("INFO", f"CSV report saved: {output_path}")
        return str(output_path)
