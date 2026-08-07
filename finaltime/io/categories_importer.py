"""Загрузка поло-возрастных категорий из Excel."""

from pathlib import Path
from typing import Optional

import openpyxl

from finaltime.core import AgeCategory, Gender, parse_date


class CategoriesImporter:
    """Загрузка категорий из Excel-файла.

    Ожидаемые колонки (можно настроить):
    - Название категории (например: "Мужчины 18-25")
    - Пол (М/Ж/Общ)
    - Мин. возраст
    - Макс. возраст
    - Общий зачёт (да/нет)
    """

    def __init__(self, logger):
        self.logger = logger

        # Индексы столбцов (0-based)
        self.col_name = 0      # Название категории
        self.col_gender = 1    # Пол
        self.col_min_age = 2   # Мин. возраст
        self.col_max_age = 3   # Макс. возраст
        self.col_open = 4      # Общий зачёт

    def load(self, file_path: str) -> list[AgeCategory]:
        """Загружает категории из Excel."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Файл категорий не найден: {file_path}")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        categories = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(c is None for c in row):
                continue

            try:
                # Название
                raw_name = row[self.col_name] if len(row) > self.col_name else None
                if raw_name is None:
                    continue
                name = str(raw_name).strip()
                if not name:
                    continue

                # Пол
                raw_gender = row[self.col_gender] if len(row) > self.col_gender else None
                gender = Gender.from_string(str(raw_gender)) if raw_gender else Gender.UNKNOWN

                # Мин. возраст
                raw_min = row[self.col_min_age] if len(row) > self.col_min_age else None
                min_age = int(raw_min) if raw_min is not None else 0

                # Макс. возраст
                raw_max = row[self.col_max_age] if len(row) > self.col_max_age else None
                max_age = int(raw_max) if raw_max is not None else 999

                # Общий зачёт
                raw_open = row[self.col_open] if len(row) > self.col_open else None
                is_open = False
                if raw_open is not None:
                    open_str = str(raw_open).strip().lower()
                    is_open = open_str in ("да", "yes", "true", "1", "+", "y")

                # Для общего зачёта устанавливаем возрастные границы широкие
                if is_open:
                    min_age = 0
                    max_age = 999

                category = AgeCategory(
                    name=name,
                    gender=gender,
                    min_age=min_age,
                    max_age=max_age,
                    is_open=is_open
                )
                categories.append(category)

            except (ValueError, IndexError) as e:
                self.logger.log("WARNING", f"Строка {row_idx}: ошибка парсинга категории - {e}")
                continue

        if not categories:
            self.logger.log("WARNING", "Не загружено ни одной категории")

        self.logger.log("INFO", f"Loaded categories: {len(categories)}")
        return categories

    def set_column_mapping(
        self,
        name: int = 0,
        gender: int = 1,
        min_age: int = 2,
        max_age: int = 3,
        open_col: int = 4
    ) -> None:
        """Настройка маппинга столбцов (0-based)."""
        self.col_name = name
        self.col_gender = gender
        self.col_min_age = min_age
        self.col_max_age = max_age
        self.col_open = open_col


def create_default_categories_file(file_path: str) -> None:
    """Создаёт пример файла категорий."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Категории"

    # Заголовки
    headers = [
        "Название", "Пол", "Мин. возраст", "Макс. возраст", "Общий зачёт"
    ]
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)

    # Примеры категорий
    categories = [
        ("Мужчины 18-25", "М", 18, 25, "Нет"),
        ("Мужчины 26-35", "М", 26, 35, "Нет"),
        ("Мужчины 36-45", "М", 36, 45, "Нет"),
        ("Мужчины 46-55", "М", 46, 55, "Нет"),
        ("Мужчины 56+", "М", 56, 999, "Нет"),
        ("Женщины 18-25", "Ж", 18, 25, "Нет"),
        ("Женщины 26-35", "Ж", 26, 35, "Нет"),
        ("Женщины 36-45", "Ж", 36, 45, "Нет"),
        ("Женщины 46-55", "Ж", 46, 55, "Нет"),
        ("Женщины 56+", "Ж", 56, 999, "Нет"),
        ("Общий зачёт мужчины", "М", 0, 999, "Да"),
        ("Общий зачёт женщины", "Ж", 0, 999, "Да"),
    ]

    for row_idx, cat in enumerate(categories, 2):
        for col_idx, value in enumerate(cat, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Ширина колонок
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14

    wb.save(file_path)