"""Экспорт отчётов в Excel и CSV."""

import csv
import openpyxl
from datetime import timedelta
from pathlib import Path
from typing import List
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter

from chronojudge.core import FinishRecord, format_time, format_time_short
from chronojudge import get_version


class Exporter:
    """Экспорт результатов в Excel и CSV."""
    
    def __init__(self, logger):
        self.logger = logger
    
    def export_excel(
        self,
        finishes: List[FinishRecord],
        output_dir: str,
        registration_file: str = ""
    ) -> str:
        """Экспорт в красивый Excel-файл."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Результаты"
        
        # Стили
        header_font = Font(name='Calibri', bold=True, size=12, color='FFFFFF')
        header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        title_font = Font(name='Calibri', bold=True, size=16, color='2C3E50')
        subtitle_font = Font(name='Calibri', size=11, color='555555')
        
        data_font = Font(name='Calibri', size=11)
        data_alignment = Alignment(horizontal='center', vertical='center')
        name_alignment = Alignment(horizontal='left', vertical='center')
        
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC'),
        )
        
        gold_fill = PatternFill(start_color='FFD700', end_color='FFD700', fill_type='solid')
        silver_fill = PatternFill(start_color='C0C0C0', end_color='C0C0C0', fill_type='solid')
        bronze_fill = PatternFill(start_color='CD7F32', end_color='CD7F32', fill_type='solid')
        
        # === ЗАГОЛОВОК ОТЧЁТА ===
        ws.merge_cells('A1:H1')
        ws['A1'] = "ПРОТОКОЛ РЕЗУЛЬТАТОВ СОРЕВНОВАНИЙ"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 35
        
        ws.merge_cells('A2:H2')
        ws['A2'] = f"Сгенерировано ChronoJudge v{get_version()}"
        ws['A2'].font = subtitle_font
        ws['A2'].alignment = Alignment(horizontal='center')
        
        if registration_file:
            ws.merge_cells('A3:H3')
            ws['A3'] = f"Файл регистрации: {Path(registration_file).name}"
            ws['A3'].font = subtitle_font
            ws['A3'].alignment = Alignment(horizontal='center')
            header_row = 5
        else:
            header_row = 4
        
        # === ЗАГОЛОВКИ ТАБЛИЦЫ ===
        headers = [
            "Место", "Номер", "ФИО участника",
            "Время таймера", "Ручное время",
            "Штраф (сек)", "Штраф (баллы)", "Итоговое время"
        ]
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        ws.row_dimensions[header_row].height = 30
        
        # === ДАННЫЕ ===
        for row_idx, record in enumerate(finishes, header_row + 1):
            place = record.place
            
            # Места 1-3 — подсветка
            row_fill = None
            if place == 1:
                row_fill = gold_fill
            elif place == 2:
                row_fill = silver_fill
            elif place == 3:
                row_fill = bronze_fill
            
            timer_str = format_time_short(record.timer_time)
            manual_str = format_time_short(record.manual_time) if record.manual_time else ""
            final_str = format_time_short(record.final_time) if record.final_time else ""
            
            row_data = [
                place,
                record.participant.number,
                record.participant.full_name,
                timer_str,
                manual_str,
                record.penalty_seconds,
                record.penalty_points,
                final_str,
            ]
            
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = data_font
                cell.border = thin_border
                
                if col_idx == 3:  # ФИО
                    cell.alignment = name_alignment
                else:
                    cell.alignment = data_alignment
                
                if row_fill:
                    cell.fill = row_fill
        
        # === НАСТРОЙКА ШИРИНЫ КОЛОНОК ===
        column_widths = [8, 10, 35, 15, 15, 12, 14, 16]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        # Заморозка верхних строк
        ws.freeze_panes = f'A{header_row + 1}'
        
        # Автофильтр
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f'A{header_row}:{last_col}{header_row + len(finishes)}'
        
        # Сохранение
        output_path = Path(output_dir) / f"Report_ChronoJudge_v{get_version()}.xlsx"
        wb.save(output_path)
        
        self.logger.log("INFO", f"Excel отчёт сохранён: {output_path}")
        return str(output_path)
    
    def export_csv(
        self,
        finishes: List[FinishRecord],
        output_dir: str
    ) -> str:
        """Экспорт в CSV."""
        output_path = Path(output_dir) / f"Report_ChronoJudge_v{get_version()}.csv"
        
        with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f, delimiter=';')
            
            # Заголовок
            writer.writerow([
                "Место", "Номер", "ФИО",
                "Время таймера", "Ручное время",
                "Штраф (сек)", "Штраф (баллы)", "Итоговое время"
            ])
            
            # Данные
            for record in finishes:
                timer_str = format_time_short(record.timer_time)
                manual_str = format_time_short(record.manual_time) if record.manual_time else ""
                final_str = format_time_short(record.final_time) if record.final_time else ""
                
                writer.writerow([
                    record.place,
                    record.participant.number,
                    record.participant.full_name,
                    timer_str,
                    manual_str,
                    record.penalty_seconds,
                    record.penalty_points,
                    final_str,
                ])
        
        self.logger.log("INFO", f"CSV отчёт сохранён: {output_path}")
        return str(output_path)